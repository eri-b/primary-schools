"""Enrich every map with official state school outcomes and finance measures.

Run after rebuilding the base JSON: python3 scripts/add_school_metrics.py
Requires openpyxl and access-parser (or mdb-export on PATH). Source downloads
and parsed, schoolwide records are cached in --source-dir, outside the repo.
Suppression is preserved; joins use state school IDs, never fuzzy school names.
"""
import argparse
import csv
import html
import http.cookiejar
import io
import json
import math
import re
import shutil
import subprocess
import urllib.parse
import urllib.request
import zipfile
from collections import defaultdict
from pathlib import Path

PUBLIC = Path(__file__).resolve().parents[1] / 'public'
SOURCES = {
    'ccd-directory.zip': 'https://nces.ed.gov/ccd/data/zip/ccd_sch_029_2425_w_1a_073025.zip',
    'ny-src.zip': 'https://data.nysed.gov/files/essa/24-25/SRC2025.zip',
    'ny-grad.zip': 'https://data.nysed.gov/files/gradrate/24-25/gradrate.zip',
    'nj-schools.xlsx': 'https://www.nj.gov/education/sprreports/download/DataFiles/2024-2025/Database_SchoolDetail.xlsx',
    'ct-data.json': 'https://data.ct.gov/resource/h28j-iix5.json?' + urllib.parse.urlencode({'$where': "schoolyear='2024-25'", '$limit': 2000}),
    'nj-spending.xlsx': 'https://www.nj.gov/education/guide/docs/2025/Total_Spending.xlsx',
}
NY_URL = 'https://data.nysed.gov/downloads.php'
NJ_URL = 'https://www.nj.gov/education/spr/download/'
CT_URL = 'https://data.ct.gov/Education/Next-Generation-Accountability-System/h28j-iix5'


def number(value):
    # No inequalities, ranges, text suppression flags, NaN, or negative sentinels.
    if value is None or isinstance(value, bool):
        return None
    try:
        result = float(str(value).replace(',', '').removeprefix('$').removesuffix('%').strip())
        return result if math.isfinite(result) and result >= 0 else None
    except ValueError:
        return None


def download(folder, name):
    path = folder / name
    if not path.exists():
        print(f'Downloading {name}', flush=True)
        temporary = path.with_suffix(path.suffix + '.part')
        request = urllib.request.Request(SOURCES[name], headers={'User-Agent': 'schools-map-data-builder/1.0'})
        with urllib.request.urlopen(request, timeout=240) as response, temporary.open('wb') as output:
            shutil.copyfileobj(response, output)
        temporary.replace(path)
    return path


def access_rows(folder, archive_name, table):
    cache = folder / ('schoolwide-' + table.replace('/', '-') + '.json')
    if cache.exists():
        return json.loads(cache.read_text())
    archive = zipfile.ZipFile(download(folder, archive_name))
    name = next(n for n in archive.namelist() if n.endswith('.mdb'))
    path = folder / name
    if not path.exists():
        archive.extract(name, folder)
    print(f'Reading NYSED {table}', flush=True)
    if shutil.which('mdb-export'):
        process = subprocess.Popen(['mdb-export', str(path), table], stdout=subprocess.PIPE, text=True)
        rows = csv.DictReader(process.stdout)
    else:
        # Pure Python fallback, useful on machines without mdbtools.
        from access_parser import AccessParser
        db = AccessParser(str(path))
        data = db.parse_table(table)
        rows = (dict(zip(data, values)) for values in zip(*data.values()))
        process = None
    result = [row for row in rows if
              row.get('SUBGROUP_NAME', 'All Students') == 'All Students'
              and (number(row.get('YEAR')) == 2025 or 'YEAR' not in row)]
    if process is not None and process.wait() != 0:
        raise RuntimeError(f'mdb-export failed for {table}')
    cache.write_text(json.dumps(result, separators=(',', ':')))
    return result


def metric(metrics, key, value, label, url, *, year='2024–25', unit='percent', note='', lower=False):
    value = number(value)
    if value is None or (unit == 'percent' and value > 100):
        return
    metrics[key] = dict(value=round(value, 1), label=label, year=year, unit=unit,
                        source=url, note=note, lowerIsBetter=lower)


def nces_directory(folder):
    with zipfile.ZipFile(download(folder, 'ccd-directory.zip')) as archive:
        name = next(n for n in archive.namelist() if n.endswith('.csv'))
        with io.TextIOWrapper(archive.open(name), encoding='utf-8-sig', newline='') as stream:
            return {row['NCESSCH']: row for row in csv.DictReader(stream)}


def state_id(row):
    return row['ST_SCHID'].split('-')[-1]


def ny_metrics(folder):
    by_beds = defaultdict(dict)
    instids = {}
    for table, key in [('ACC EM Chronic Absenteeism', 'absenteeism-em'), ('ACC HS Chronic Absenteeism', 'absenteeism-hs')]:
        for row in access_rows(folder, 'ny-src.zip', table):
            beds = row['ENTITY_CD']
            instids[beds] = row['INSTITUTION_ID']
            if row.get('DATA_REP_FLAG') != 'Y' or row.get('PARTIAL_DATA_FLAG') == 'Y':
                continue
            metric(by_beds[beds], key, row['ABSENT_RATE'],
                   'Chronic absenteeism (grades 1–8)' if key.endswith('em') else 'Chronic absenteeism (grades 9–12)',
                   NY_URL, lower=True, note='Students absent at least 10% of enrolled instructional days. Ungraded age-equivalent students are included.')
    for row in access_rows(folder, 'ny-src.zip', 'Expenditures per Pupil'):
        beds = row['ENTITY_CD']
        instids[beds] = row['INSTITUTION_ID']
        if row.get('DATA_REPORTED_EXP') != 'Y':
            continue
        metric(by_beds[beds], 'spending', row['PER_FED_STATE_LOCAL_EXP'], 'Spending per pupil', NY_URL, unit='currency',
               note='School-level federal, state, and local expenditures per pupil, including allocated district costs. Spending is a resource measure, not a quality rating.')
    for row in access_rows(folder, 'ny-src.zip', 'Annual Regents Exams'):
        beds = row['ENTITY_CD']
        if row.get('INSTITUTION_ID'):
            instids[beds] = row['INSTITUTION_ID']
        if number(row.get('TESTED')) is None:
            continue
        subject = row['SUBJECT']
        key = 'regents-' + re.sub('[^a-z0-9]+', '-', subject.lower()).strip('-')
        metric(by_beds[beds], key, row['PER_PROF'], subject + ' pass rate', NY_URL,
               note=f"Annual exam results; {int(number(row['TESTED'])):,} tested. Pass rate is the published proficiency percentage, not a diploma rate. Includes all grades taking this exam.")
    # Aggregate only published grade counts; never recover a suppressed count.
    for subject in ('ELA', 'MATH'):
        grouped = defaultdict(list)
        for row in access_rows(folder, 'ny-src.zip', 'Annual EM ' + subject):
            grouped[row['ENTITY_CD']].append(row)
            if row.get('INSTITUTION_ID'):
                instids[row['ENTITY_CD']] = row['INSTITUTION_ID']
        for beds, rows in grouped.items():
            for level, grades in [('elementary', (3, 4, 5)), ('middle', (6, 7, 8))]:
                selected = [r for r in rows if r['ASSESSMENT_NAME'] in ({f'Combined{g}Math' for g in grades} if subject == 'MATH' and level == 'middle' else {f'{subject}{g}' for g in grades})]
                counts = [(number(r['NUM_TESTED']), number(r['NUM_PROF'])) for r in selected]
                if not counts or any(t is None or p is None for t, p in counts):
                    continue
                tested = sum(t for t, _ in counts)
                if tested:
                    metric(by_beds[beds], subject.lower() + '-' + level, 100 * sum(p for _, p in counts) / tested,
                           f'{subject.title() if subject == "MATH" else subject} proficiency (grades {grades[0]}–{grades[-1]})', NY_URL,
                           note=f'Weighted by {int(tested):,} published tested counts. NYSTP proficiency; grades 6–8 math includes accelerated Regents math via NYSED combined-grade counts. Schools with suppressed grade counts are unavailable.')
    for row in access_rows(folder, 'ny-grad.zip', 'GRAD_RATE_AND_OUTCOMES_2025'):
        if row.get('AGGREGATION_TYPE') != 'School' or row.get('MEMBERSHIP_DESC') != '2021 Total Cohort - 4 Year Outcome - August 2025':
            continue
        beds = row['AGGREGATION_CODE']
        instids[beds] = row['INSTITUTION_ID']
        cohort = number(row['ENROLL_CNT'])
        note = '2021 entering grade-9 cohort, four-year outcome through August 2025. Percent of the entire cohort, not just graduates. Rates use published counts where available.'
        for key, field, label, lower in [
            ('graduation', 'GRAD', 'Four-year graduation rate', False),
            ('advanced-regents', 'REG_ADV', 'Advanced Regents diploma rate', False),
            ('dropout', 'DROPOUT', 'Four-year cohort dropout rate', True),
        ]:
            count = number(row[field + '_CNT'])
            rate = 100 * count / cohort if count is not None and cohort else number(row[field + '_PCT'])
            metric(by_beds[beds], key, rate, label, NY_URL, lower=lower, note=note)
        standard, advanced = number(row['REG_CNT']), number(row['REG_ADV_CNT'])
        if cohort and standard is not None and advanced is not None:
            metric(by_beds[beds], 'regents-diploma', 100 * (standard + advanced) / cohort,
                   'Regents diploma rate (including advanced)', NY_URL, note=note)
    return by_beds, instids


def nj_key(row):
    return (str(row['CountyCode']).zfill(2), str(row['DistrictCode']).zfill(4), str(row['SchoolCode']).zfill(3))


def nj_rows(workbook, name):
    rows = iter(workbook[name].values)
    for _ in range(3):
        next(rows)
    headers = next(rows)
    for values in rows:
        if values[0] is not None:
            yield dict(zip(headers, values))


def nj_metrics(folder):
    cache = folder / 'nj-schoolwide.json'
    tables = json.loads(cache.read_text()) if cache.exists() else {}
    names = ['ELAPerformanceByTest', 'MathPerformancebyTest', 'GraduationRateTrends', 'DropoutRateTrends', 'ChronicAbsenteeismTrends', 'NJGPA']
    missing = [name for name in names if name not in tables]
    if missing:
        import openpyxl
        workbook = openpyxl.load_workbook(download(folder, 'nj-schools.xlsx'), read_only=True, data_only=True)
        for name in missing:
            print('Reading NJDOE ' + name, flush=True)
            tables[name] = [r for r in nj_rows(workbook, name) if r.get('SchoolYear') == '2024-25' and r.get('StudentGroup', 'All Students') == 'All Students']
        cache.write_text(json.dumps(tables, separators=(',', ':')))
    result = defaultdict(dict)
    for row in tables['ChronicAbsenteeismTrends']:
        metric(result[nj_key(row)], 'absenteeism', row['ChronicAbsenteeismRate_School'], 'Chronic absenteeism (K–12)', NJ_URL,
               lower=True, note='Schoolwide students absent at least 10% of enrolled days; includes grades outside the selected school level.')
    for row in tables['GraduationRateTrends']:
        metric(result[nj_key(row)], 'graduation', row['4YrGraduationRateFederal_School'], 'Four-year graduation rate (federal)', NJ_URL,
               note='Federal adjusted cohort graduation rate. Cohort: ' + str(row['4YrCohort']))
    for row in tables['DropoutRateTrends']:
        metric(result[nj_key(row)], 'dropout', row['Dropout_School'], 'Annual dropout rate (grades 9–12)', NJ_URL,
               lower=True, note='Students in grades 9–12 dropping out during the school year; not a four-year cohort dropout rate.')
    for row in tables['NJGPA']:
        for subject in ('ELA', 'Math'):
            metric(result[nj_key(row)], 'njgpa-' + subject.lower(), row[subject + 'GraduationReadyPct_School'],
                   'NJGPA ' + subject + ' graduation readiness', NJ_URL,
                   note='Percentage meeting the NJ Graduation Proficiency Assessment graduation-ready threshold, normally taken in grade 11. This is an assessment result, not a graduation rate.')
    for table, subject in [('ELAPerformanceByTest', 'ELA'), ('MathPerformancebyTest', 'Math')]:
        grouped = defaultdict(list)
        for row in tables[table]:
            grouped[nj_key(row)].append(row)
        for sid, rows in grouped.items():
            grade_groups = [('elementary', (3, 4, 5)), ('middle', (6, 7, 8))]
            if subject == 'ELA':
                grade_groups.append(('high', (9,)))
            for level, grades in grade_groups:
                selected = [r for r in rows if r['GradeTest'] in {f'Grade {g}' for g in grades}]
                counts = [(number(r['ValidScores_School']), number(r['ProficiencyRate_School'])) for r in selected]
                if not counts or any(t is None or p is None for t, p in counts):
                    continue
                tested = sum(t for t, _ in counts)
                if tested:
                    metric(result[sid], subject.lower() + '-' + level, sum(t*p for t, p in counts) / tested,
                           f'{subject} proficiency (grade {grades[0]})' if len(grades) == 1 else f'{subject} proficiency (grades {grades[0]}–{grades[-1]})', NJ_URL,
                           note=f'NJSLA Levels 4–5, weighted by {int(tested):,} published valid scores. Accelerated course exams are reported separately. Suppressed grades remain unavailable.')
            for row in rows:
                if row['GradeTest'].startswith('Grade '):
                    continue
                key = 'njsla-' + re.sub('[^a-z0-9]+', '-', row['GradeTest'].lower()).strip('-')
                metric(result[sid], key, row['ProficiencyRate_School'], 'NJSLA ' + row['GradeTest'] + ' proficiency', NJ_URL,
                       note='Published course exam proficiency (Levels 4–5), across all grades taking the exam.')
    return result


def nj_spending(folder):
    import openpyxl
    workbook = openpyxl.load_workbook(download(folder, 'nj-spending.xlsx'), read_only=True, data_only=True)
    # NJ's 2025 Taxpayers' Guide reports audited 2023–24 actual spending.
    # District codes are unique statewide; never match by district name.
    result = {}
    for values in workbook.active.iter_rows(min_row=3, values_only=True):
        code = number(values[1])
        if code is not None:
            result[str(int(code)).zfill(4)] = number(values[9])
    return result


def ct_report(folder, filename, program, parameters):
    """Download a public EdSight export through its published Guest option."""
    path = folder / filename
    if path.exists():
        return path
    opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))
    url = 'https://edsight.ct.gov/SASStoredProcess/guest?' + urllib.parse.urlencode({
        '_program': '/CTDOE/EdSight/Release/Reporting/Public/Reports/StoredProcesses/' + program,
        **parameters,
    })
    response = opener.open(url, timeout=120)
    content = response.read()
    if b'Logon Manager' in content:
        execution = re.search(r'name="execution" value="([^"]+)"', content.decode()).group(1)
        target = urllib.parse.urljoin(response.url, 'login')
        response = opener.open(target, data=urllib.parse.urlencode({
            'execution': html.unescape(execution), '_eventId': 'guest', 'fallback': 'true',
        }).encode(), timeout=120)
        content = response.read()
    if b'Logon Manager' in content or b'<html' in content.lower():
        raise RuntimeError(f'EdSight did not return a CSV export for {filename}')
    path.write_bytes(content)
    return path


def ct_spending(folder):
    path = ct_report(folder, 'ct-spending.csv', 'EFSSchoolLevelbyFunctionExport', {
        '_year': '2024-25', '_district': 'All Districts', '_school': 'All Schools',
    })
    with path.open(encoding='utf-8-sig', newline='') as stream:
        rows = csv.reader(stream)
        headers = next(row for row in rows if row and row[0] == 'District')
        result = {}
        for values in rows:
            row = dict(zip(headers, values))
            # EdSight writes Excel text literals, e.g. ="2850113". Treat as text.
            match = re.fullmatch(r'="(\d{7})"', row.get('School Code', ''))
            if match:
                result[match.group(1)] = number(row.get('Total'))
        return result


def ct_graduation(folder):
    path = ct_report(folder, 'ct-grad.csv', 'GraduationExport', {
        '_year': '2024-25', '_district': 'All Districts', '_school': 'All Schools',
        '_subgroup': 'All Students', '_rate': ' ', '_gradcat': 'grad',
    })
    with path.open(encoding='utf-8-sig', newline='') as stream:
        rows = csv.reader(stream)
        headers = next(row for row in rows if row and row[0] == 'District')
        result = {}
        for values in rows:
            row = dict(zip(headers, values))
            match = re.fullmatch(r'="(\d{7})"', row.get('School Code', ''))
            if match:
                result[match.group(1)] = number(row.get('Graduation Rate'))
        return result


def ct_metrics(folder):
    rows = json.loads(download(folder, 'ct-data.json').read_text())
    if len(rows) >= 2000:
        raise RuntimeError('Connecticut response may be truncated')
    result = defaultdict(dict)
    spending = ct_spending(folder)
    graduation = ct_graduation(folder)
    for row in rows:
        if row['schoolcode'] == '0000000':
            continue
        metrics = result[row['schoolcode']]
        metric(metrics, 'spending', spending.get(row['schoolcode']), 'School-level spending per pupil',
               'https://public-edsight.ct.gov/overview/per-pupil-expenditures-by-function---district/per-pupil-expenditures-by-function---school', unit='currency',
               note='School-level instructional and support expenditures. Excludes non-local food services, debt, capital other than equipment, adult education, community services, and state teacher-retirement contributions. Not comparable to other states’ spending definitions.')
        metric(metrics, 'graduation', graduation.get(row['schoolcode']), 'Four-year graduation rate',
               'https://public-edsight.ct.gov/performance/four-year-graduation-rates',
               note='2024–25 four-year cohort graduation rate, including early and summer graduates receiving a standard diploma. Missing or suppressed cohort results remain unavailable.')
        for key, field, label, unit, factor, lower in [
            ('ela-index', 'ind1ela_all_rate', 'ELA performance index', 'index', 1, False),
            ('math-index', 'ind1math_all_rate', 'Math performance index', 'index', 1, False),
            ('absenteeism', 'ind4rate', 'Chronic absenteeism', 'percent', 100, True),
            ('ela-growth', 'ind2ela_all_rate', 'ELA growth target achieved', 'percent', 100, False),
            ('math-growth', 'ind2math_all_rate', 'Math growth target achieved', 'percent', 100, False),
            ('college-coursework', 'ind5rate', 'Postsecondary preparation coursework', 'percent', 100, False),
            ('college-readiness', 'ind6rate', 'Postsecondary readiness', 'percent', 100, False),
        ]:
            value = number(row.get(field))
            metric(metrics, key, value * factor if value is not None else None, label, CT_URL, unit=unit, lower=lower,
                   note='CSDE 2024–25 accountability indicator; may use an earlier graduating cohort. Schoolwide, across applicable grades. Performance index is a 0–100 achievement index, not a proficiency percentage.' if unit == 'index' else 'CSDE 2024–25 accountability indicator. Graduation uses the cohort specified by CSDE; other indicators cover applicable grades schoolwide.')
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-dir', type=Path, default=Path('/tmp/school-metrics'))
    parser.add_argument('--states', nargs='+', choices=['ny', 'nj', 'ct'], default=['ny', 'nj', 'ct'])
    args = parser.parse_args()
    args.source_dir.mkdir(parents=True, exist_ok=True)
    directory = nces_directory(args.source_dir)
    ny, instids = ny_metrics(args.source_dir) if 'ny' in args.states else ({}, {})
    nj = nj_metrics(args.source_dir) if 'nj' in args.states else {}
    nj_finance = nj_spending(args.source_dir) if 'nj' in args.states else {}
    ct = ct_metrics(args.source_dir) if 'ct' in args.states else {}
    dbn_to_beds = {}
    for row in directory.values():
        beds = state_id(row)
        if row['ST_SCHID'].startswith('NY-') and len(beds) == 12:
            # CCD district number is contained in the BEDS identifier for NYC.
            county_letter = {'31': 'M', '32': 'X', '33': 'K', '34': 'Q', '35': 'R'}.get(beds[:2])
            if county_letter and 1 <= int(beds[2:4]) <= 32:
                dbn_to_beds[beds[2:4] + county_letter + beds[-3:]] = beds
    for path in sorted(PUBLIC.glob('*schools.json')):
        region = path.stem.removeprefix('middle-').removeprefix('high-').removesuffix('-schools')
        state = 'nj' if region == 'new-jersey' else 'ct' if region == 'connecticut' else 'ny'
        if state not in args.states:
            continue
        schools = json.loads(path.read_text())
        level = 'middle' if path.name.startswith('middle-') else 'high' if path.name.startswith('high-') else 'elementary'
        matched = 0
        for school in schools:
            row = directory.get(school.get('id', ''))
            if row:
                sid = state_id(row)
                if state == 'ny':
                    metrics = dict(ny.get(sid, {}))
                elif state == 'ct':
                    metrics = dict(ct.get(sid, {}))
                else:
                    district = row['ST_LEAID'].removeprefix('NJ-')
                    metrics = dict(nj.get((district[:2], district[2:].zfill(4), sid.zfill(3)), {}))
                    metric(metrics, 'district-spending', nj_finance.get(district[2:].zfill(4)), 'District spending per pupil',
                           'https://www.nj.gov/education/guide/2025tges.shtml', unit='currency', year='2023–24',
                           note='Audited actual district total spending per pupil from NJDOE’s 2025 Taxpayers’ Guide. Same district value for every school; not school-specific. Includes district-paid costs for students sent elsewhere.')
            elif school.get('dbn'):
                sid = dbn_to_beds.get(school['dbn'])
                metrics = dict(ny.get(sid, {}))
            else:
                metrics = {}
            for key in list(metrics):
                if key.endswith(('-elementary', '-middle', '-high')) and not key.endswith('-' + level):
                    del metrics[key]
                elif key.startswith('absenteeism-') and key != ('absenteeism-hs' if level == 'high' else 'absenteeism-em'):
                    del metrics[key]
                elif level != 'high' and key in ('graduation', 'dropout', 'advanced-regents', 'regents-diploma', 'college-coursework', 'college-readiness'):
                    del metrics[key]
                elif level != 'high' and key.startswith('njgpa-'):
                    del metrics[key]
            school['metrics'] = metrics
            matched += bool(metrics)
        path.write_text(json.dumps(schools, separators=(',', ':')) + '\n')
        print(f'{path.name}: {matched}/{len(schools)} schools with additional measures', flush=True)


if __name__ == '__main__':
    main()
